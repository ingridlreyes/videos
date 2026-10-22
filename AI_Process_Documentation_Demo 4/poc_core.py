"""PoC offline: document UI procedures from supplied snapshots or review candidate video frames.

Important: extraction of visual changes does NOT establish which buttons were clicked.
Speaker filtering requires a diarized transcript upstream.
"""
import os, json, base64, math, subprocess, uuid
from pathlib import Path
from io import BytesIO
from typing import Optional
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
import cv2

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'outputs'
OUT.mkdir(parents=True, exist_ok=True)
NAVY = RGBColor(20, 37, 66)
BLUE = RGBColor(36, 111, 211)
MINT = RGBColor(25, 137, 134)
DARK = (20, 37, 66)
GRAY = RGBColor(100, 115, 133)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(241, 246, 251)


def get_font(size=20, bold=False):
    for f in (["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"] if bold else ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]):
        try: return ImageFont.truetype(f, size)
        except Exception: pass
    return ImageFont.load_default()


def fit_text(box, text, size=20, color=WHITE, bold=False):
    box.text = str(text)
    for p in box.text_frame.paragraphs:
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    box.text_frame.word_wrap = True


def add_shape(slide, left, top, width, height, fill, radius=False, stroke=None):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    obj = slide.shapes.add_shape(typ, Inches(left), Inches(top), Inches(width), Inches(height))
    obj.fill.solid(); obj.fill.fore_color.rgb = fill
    if stroke:
        obj.line.color.rgb = stroke; obj.line.width = Pt(1.4)
    else: obj.line.fill.background()
    return obj


def add_text(slide, text, x, y, w, h, size=20, color=NAVY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.margin_left=Inches(0.04)
    box.text_frame.margin_top=Inches(0.01)
    fit_text(box, text, size=size, color=color, bold=bold)
    return box


def hms(sec):
    sec = int(float(sec)); return f'{sec//3600:02d}:{(sec%3600)//60:02d}:{sec%60:02d}'


def validate_bbox(bbox):
    if bbox is None: return None
    try:
        vals = [float(x) for x in bbox]
        if len(vals)!=4 or any(not math.isfinite(v) for v in vals): return None
        x,y,w,h=vals
        if min(vals)<0 or max(vals)>1 or w<=0 or h<=0 or x+w>1.001 or y+h>1.001: return None
        return vals
    except Exception: return None


def annotate(image_path, bbox, label, out_path):
    im = Image.open(image_path).convert('RGB'); dr=ImageDraw.Draw(im)
    bbox=validate_bbox(bbox)
    if bbox:
        w,h=im.size; x,y,bw,bh=bbox
        xy=(int(x*w), int(y*h), int((x+bw)*w), int((y+bh)*h))
        lw=max(4,int(w/250))
        dr.rounded_rectangle(xy, radius=9, outline=(234,62,48), width=lw)
        cy=max(0,int(y*h)-43)
        dr.rounded_rectangle((int(x*w),cy,int(x*w)+min(390,w-int(x*w)),cy+37),radius=6,fill=(234,62,48))
        dr.text((int(x*w)+12,cy+6), label[:25],fill=(255,255,255),font=get_font(max(12,int(w/90)),True))
    im.save(out_path,quality=94)
    return str(out_path)


def source_image(data, key='image'):
    raw=Path(data[key]); return raw if raw.is_absolute() else ROOT/raw


def load_process(source=None):
    p=Path(source) if source else ROOT/'example_steps.json'
    dat=json.loads(p.read_text(encoding='utf-8'))
    for st in dat['steps']:
        st['bbox']=validate_bbox(st.get('bbox'))
        if st['bbox'] is None: st['review_required']=True
    return dat


def presentation(data, filename='Demo_Flujo_Anotado.pptx'):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    sl=prs.slides.add_slide(blank)
    add_shape(sl,0,0,13.333,7.5,NAVY)
    add_shape(sl,0.8,1.08,1.18,.10,RGBColor(86,210,190))
    add_text(sl,'AI PROCESS DOCUMENTATION',.8,1.36,11,0.5,16,RGBColor(139,209,215),True)
    add_text(sl,data.get('title','Procedimiento documentado'),.8,2.10,11.2,1.55,36,WHITE,True)
    subtitle='DEMO SINTÉTICA' if data.get('source_type') in (None,'demo') else 'BORRADOR PARA REVISIÓN'
    add_text(sl,subtitle+' • Evidencias y pasos revisables',.8,4.66,11,0.8,20,RGBColor(220,231,242))
    add_text(sl,'Revisar y aprobar pasos, capturas y privacidad antes de publicar.',.8,6.72,11,.4,12,RGBColor(201,216,234))
    sl=prs.slides.add_slide(blank);add_shape(sl,0,0,13.333,7.5,PALE)
    add_text(sl,'Mapa del proceso',.75,.65,11,0.7,30,NAVY,True)
    add_text(sl,'Flujo reconstruido para convertir una sesión en documentación accionable',.75,1.4,12,.55,17,GRAY)
    steps=data['steps']
    count=len(steps); each=min(2.30, (11.8-(count-1)*.24)/max(1,count)); x=.76
    for i,st in enumerate(steps):
        if x+each>12.8: break
        add_shape(sl,x,2.55,each,2.05,WHITE,True,RGBColor(215,227,237))
        add_shape(sl,x+.13,2.78,.46,.46,BLUE,True)
        add_text(sl,str(i+1),x+.22,2.84,.29,.27,13,WHITE,True)
        add_text(sl,st['title'],x+.15,3.39,each-.30,.90,15,NAVY,True)
        x+=each+.24
    add_text(sl,'Reglas aplicables: '+data.get('rules_summary','Incluir sólo acciones instruccionales.'),.75,6.2,11.7,.7,16,GRAY)
    for i,st in enumerate(steps,1):
        sl=prs.slides.add_slide(blank);add_shape(sl,0,0,13.333,7.5,WHITE)
        add_shape(sl,0,0,13.333,.16,BLUE)
        add_text(sl,f'PASO {i:02d} / {len(steps):02d}',.55,.42,4,.4,13,BLUE,True)
        add_text(sl,st['title'],.55,.98,12.0,.76,28,NAVY,True)
        im=source_image(st)
        ann=OUT/f'captura_paso_{i:02d}.png'; annotate(im,st.get('bbox'),st.get('target_label','Acción'),ann)
        iw,ih=Image.open(im).size
        alloc_w,alloc_h=8.2,4.63; scale=min(alloc_w/(iw/ih),alloc_h)
        width=scale*(iw/ih); height=scale
        x=.55+(alloc_w-width)/2; y=2.03+(alloc_h-height)/2
        # PowerPoint uses original screenshot; callout rectangle is a NATIVE, editable shape.
        sl.shapes.add_picture(str(im),Inches(x),Inches(y),width=Inches(width),height=Inches(height))
        bbox=validate_bbox(st.get('bbox'))
        if bbox:
            bx,by,bw,bh=bbox
            rect=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+bx*width),Inches(y+by*height),
                                     Inches(bw*width),Inches(bh*height))
            rect.fill.background();rect.line.color.rgb=RGBColor(234,62,48);rect.line.width=Pt(3)
            tag=add_shape(sl,x+bx*width,max(y,y+by*height-.33),
                          min(2.2,max(.9,width-bx*width)),.29,RGBColor(234,62,48),True)
            add_text(sl,str(st.get('target_label','Control'))[:25],x+bx*width+.05,
                     max(y,y+by*height-.33)+.035,min(2.05,max(.75,width-bx*width-.1)),.22,
                     10,WHITE,True)
        add_shape(sl,9.04,2.03,3.7,4.61,PALE,True)
        add_text(sl,'INSTRUCCIÓN',9.35,2.35,2.9,.38,12,BLUE,True)
        add_text(sl,st['instruction'],9.35,2.91,3.0,1.72,19,NAVY,True)
        add_text(sl,'Referencia en video',9.35,5.1,3.0,.40,13,GRAY)
        add_text(sl,hms(st.get('second',0)),9.35,5.50,3.1,.40,16,BLUE,True)
        status='Pendiente de revisión' if st.get('review_required',False) else 'Ejemplo anotado'
        add_text(sl,status,.55,6.93,8.4,.35,12,GRAY)
    sl=prs.slides.add_slide(blank); add_shape(sl,0,0,13.333,7.5,NAVY)
    add_text(sl,'Validación antes de publicar',.8,.85,11,1,30,WHITE,True)
    items=['¿La secuencia coincide con lo que realmente hizo el instructor?',
           '¿Las marcas señalan el elemento correcto?',
           '¿Se eliminaron interrupciones, opiniones y preguntas del flujo?',
           '¿Se protegieron datos personales, credenciales y pantallas sensibles?',
           '¿Los pasos tienen referencia al minuto del video original?']
    for k,item in enumerate(items):
        add_shape(sl,.83,2.0+k*.85,.22,.22,RGBColor(89,208,189),True)
        add_text(sl,item,1.28,1.91+k*.85,11,.63,17,WHITE)
    out=OUT/filename; prs.save(out); return str(out)


def guide(data,filename='Guia_Usuario_Demo.docx'):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=DInches(.76);sec.bottom_margin=DInches(.68)
    sec.left_margin=DInches(.80);sec.right_margin=DInches(.80)
    doc.styles['Normal'].font.name='Aptos'
    doc.styles['Normal'].font.size=DPt(10)
    doc.styles['Normal'].paragraph_format.space_after=DPt(7)
    heading=doc.styles['Heading 1'];heading.font.name='Aptos Display';heading.font.size=DPt(17);heading.font.color.rgb=DRGB(20,37,66)
    doc.add_paragraph('GUÍA DE USUARIO  |  '+('DEMO SINTÉTICA' if data.get('source_type') in (None,'demo') else 'BORRADOR EN REVISIÓN'),style='Subtitle')
    doc.add_heading(data.get('title','Procedimiento'),0)
    doc.add_paragraph('Documento por validar. Las instrucciones y señalizaciones requieren revisión humana antes de su uso operativo.')
    doc.add_heading('Objetivo',1)
    doc.add_paragraph(data.get('goal','Ejecutar el procedimiento descrito mediante instrucciones visuales y secuenciales.'))
    doc.add_heading('Antes de comenzar',1)
    doc.add_paragraph(data.get('prerequisites','Disponer de acceso autorizado al sistema de ejemplo.'))
    doc.add_heading('Procedimiento',1)
    for i,st in enumerate(data['steps'],1):
        heading_step=doc.add_heading(f'Paso {i}. {st["title"]}',2)
        heading_step.paragraph_format.keep_with_next=True
        instruction_p=doc.add_paragraph(st['instruction'])
        instruction_p.paragraph_format.keep_with_next=True
        if st.get('detail'):
            details_p=doc.add_paragraph(st['detail'])
            details_p.paragraph_format.keep_with_next=True
        pic=OUT/f'captura_paso_{i:02d}.png'
        if not pic.exists(): annotate(source_image(st),st.get('bbox'),st.get('target_label','Acción'),pic)
        fig_p=doc.add_paragraph()
        fig_p.add_run().add_picture(str(pic),width=DInches(5.2))
        fig_p.paragraph_format.keep_with_next=True
        p=doc.add_paragraph(f'Evidencia temporal en el video: {hms(st.get("second",0))}.')
        for r in p.runs:r.italic=True;r.font.size=DPt(9)
    doc.add_heading('Reglas utilizadas',1)
    for rule in data.get('rules',[]): doc.add_paragraph(rule,style='List Bullet')
    doc.add_heading('Dudas y aclaraciones (fuera del procedimiento)',1)
    doc.add_paragraph('Las preguntas y respuestas pueden documentarse separadamente tras su clasificación y validación. Esta sección no implica que se hayan detectado preguntas automáticamente.')
    doc.add_heading('Control de calidad',1)
    doc.add_paragraph('Revisar autoría, permisos, capturas, ubicación de marcadores y el contenido de cada instrucción. Los clics no se verifican a partir de un único fotograma.')
    footer=sec.footer.paragraphs[0];footer.text='AI PROCESS DOCUMENTATION • Revisión requerida • Fuente: '+str(data.get('source_type','demo'))
    footer.alignment=2
    out=OUT/filename;doc.save(out);return str(out)


def generate(data, suffix='Demo'):
    return presentation(data,f'{suffix}_Flujo_Anotado.pptx'), guide(data,f'{suffix}_Guia_Usuario.docx')


def extract_candidate_frames(video_path, every_sec=3, max_seconds=600, max_candidates=24, diff_threshold=.018, max_gap_seconds=9):
    """Select changed screenshots; no click/semantic understanding. Bound runtime for local PoC.
    Production: streaming, backpressure, job queue and resumable processing.
    """
    cap=cv2.VideoCapture(str(video_path)); fps=cap.get(cv2.CAP_PROP_FPS) or 25
    n=cap.get(cv2.CAP_PROP_FRAME_COUNT); dur=n/fps if n>0 else max_seconds
    end_sec=min(float(dur),float(max_seconds))
    prev=None; found=[]; t=0; last_selected=-1e9
    saved_dir=OUT/'video_candidates'/uuid.uuid4().hex;saved_dir.mkdir(parents=True,exist_ok=True)
    while t<end_sec and len(found)<max_candidates:
        cap.set(cv2.CAP_PROP_POS_MSEC,t*1000);ok,frame=cap.read()
        if not ok: break
        gray=cv2.cvtColor(cv2.resize(frame,(320,180)),cv2.COLOR_BGR2GRAY)
        delta=1 if prev is None else float(cv2.absdiff(gray,prev).mean()/255)
        prev=gray
        if delta>=diff_threshold or not found or (t-last_selected>=max_gap_seconds):
            path=saved_dir/f'candidate_{len(found)+1:03d}.jpg';cv2.imwrite(str(path),frame,[int(cv2.IMWRITE_JPEG_QUALITY),92])
            last_selected=t
            found.append({'time':hms(t),'second':round(t,1),'image':str(path),'difference':round(delta,4),'title':f'Captura candidata {len(found)+1}',
                         'instruction':'Revisar esta pantalla; no se ha confirmado ninguna interacción.',
                         'target_label':'Revisar','bbox':None,'review_required':True})
        t+=max(1,int(every_sec))
    cap.release();return found,float(dur)


def filter_diarized_transcript(segments, main_speaker='0', keep_questions=False):
    """Requires upstream diarization; a speaker number alone is not verified identity."""
    chosen=[];qa=[]
    for s in segments:
        if str(s.get('speaker'))==str(main_speaker):chosen.append(s)
        elif keep_questions:qa.append(s)
    return chosen,qa


def analyze_screen_azure(image_path, transcript_excerpt='', rules=''):
    """Optional Azure OpenAI v1 endpoint. A visual guess is not a verified click."""
    from openai import OpenAI
    endpoint=os.getenv('AZURE_OPENAI_BASE_URL','').rstrip('/')+'/'
    model=os.getenv('AZURE_OPENAI_DEPLOYMENT')
    key=os.getenv('AZURE_OPENAI_API_KEY')
    if not (key and model and endpoint.startswith('https://')):
        raise RuntimeError('Faltan AZURE_OPENAI_BASE_URL, AZURE_OPENAI_API_KEY o AZURE_OPENAI_DEPLOYMENT.')
    enc=base64.b64encode(Path(image_path).read_bytes()).decode('ascii')
    client=OpenAI(api_key=key,base_url=endpoint)
    system=('Eres analista de procedimientos digitales. Solo usa evidencias visuales y transcripcion provistas. '
            'Un fotograma NO demuestra que hubo un clic: redacta como accion propuesta si no hay confirmacion. '
            'No inventes opciones o elementos, no deduzcas passwords ni ignores permisos. '
            'Responde JSON valido con llaves title,instruction,target_label,bbox,confidence,review_required. '
            'bbox es [x,y,width,height] normalizado 0..1; usa null si no identificas el control. '
            'Si hay ambiguedad, review_required=true. Sigue las reglas sin violar estas restricciones.')
    msg=f'Transcripcion principal cercana: {transcript_excerpt or "No disponible"}. Reglas: {rules or "ninguna"}.'
    resp=client.chat.completions.create(model=model,messages=[{'role':'system','content':system},
        {'role':'user','content':[{'type':'text','text':msg},
        {'type':'image_url','image_url':{'url':f'data:image/jpeg;base64,{enc}','detail':'high'}}]}],
        response_format={'type':'json_object'},max_completion_tokens=650)
    obj=json.loads(resp.choices[0].message.content)
    bbox=validate_bbox(obj.get('bbox'))
    return {'title':obj.get('title','Accion por revisar'), 'instruction':obj.get('instruction','Revisar la pantalla'),
            'target_label':obj.get('target_label','Elemento'), 'bbox':bbox,
            'confidence':obj.get('confidence','no especificada'),
            'review_required':bool(obj.get('review_required',True)) or bbox is None}
